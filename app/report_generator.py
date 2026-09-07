import io
import csv
import json
from typing import Dict, Any

from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.units import inch

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

class ReportGenerator:
    
    @staticmethod
    def generate_json(data: Dict[str, Any]) -> str:
        """Returns pretty-printed JSON string."""
        return json.dumps(data, indent=2, ensure_ascii=False)

    @staticmethod
    def generate_markdown(data: Dict[str, Any]) -> str:
        """Generates comprehensive 8-Stage DSP Markdown report."""
        filename = data.get("filename", "Unknown")
        file_id = data.get("file_id", "N/A")
        total_lat = data.get("total_latency_ms", 0)
        
        filt = data.get("filtering_stats", {})
        fft = data.get("fft_features", {})
        freq = data.get("frequency_estimation", {}) or data.get("parameters", {})
        sync = data.get("synchronization", {})
        amc = data.get("amc", {})
        noise = data.get("noise_estimation", {}) or data.get("preprocessing_stats", {})
        demod = data.get("demodulation", {}) or data.get("symbols_to_bits", {})
        sig_class = data.get("signal_classification", {})
        telem = data.get("telemetry", {})
        framing = data.get("framing", {})
        crc = data.get("crc", {})
        fec = data.get("fec", {})
        descramble = data.get("descrambler", {})
        deinter = data.get("deinterleaver", {})
        latencies = data.get("latencies", {})

        md = []
        md.append("# AetherSpectra // Autonomous Signal Intelligence & Telemetry Dossier")
        md.append(f"**Target Ingest File:** `{filename}` | **File ID:** `{file_id}` | **Total Latency:** `{total_lat} ms`\n")
        md.append("---\n")

        md.append("## 1. Ordered 8-Stage DSP Pipeline Execution Summary")
        md.append("| Stage # | Stage Name | Output / Key Metric | Execution Latency |")
        md.append("| :--- | :--- | :--- | :--- |")
        md.append(f"| **1** | **Filtering** | {filt.get('filter_type', 'Butterworth Bandpass')} (DC I: {filt.get('dc_offset_i', 0):.4f}) | `{latencies.get('filtering', 0)} ms` |")
        md.append(f"| **2** | **FFT Analysis** | Peak: {fft.get('peak_power_db', 0):.1f} dB (SFM: {fft.get('spectral_flatness', 0):.3f}) | `{latencies.get('fft', 0)} ms` |")
        md.append(f"| **3** | **Frequency Estimation** | fc: {freq.get('carrier_freq_hz', 0)/1e6:.4f} MHz (BW: {freq.get('occupied_bandwidth_99_hz', freq.get('occupied_bw_99_hz', 0))/1e3:.1f} kHz) | `{latencies.get('frequency_estimation', 0)} ms` |")
        md.append(f"| **4** | **Synchronization** | Costas: {sync.get('carrier_lock_status', 'ACQUIRING')} (EVM: {sync.get('evm_db', 0):.1f} dB) | `{latencies.get('synchronization', 0)} ms` |")
        md.append(f"| **5** | **Modulation Detection** | **{amc.get('identified_modulation', 'N/A')}** ({round(amc.get('confidence', 0)*100, 1)}%) | `{latencies.get('modulation_detection', 0)} ms` |")
        md.append(f"| **6** | **Noise Estimation** | SNR: {noise.get('snr_db', noise.get('snr_m2m4_db', 0)):.1f} dB (N0: {noise.get('noise_floor_n0_dbm_hz', -120):.1f} dBm/Hz) | `{latencies.get('noise_estimation', 0)} ms` |")
        md.append(f"| **7** | **Demodulation** | {demod.get('bit_count', 0)} Demodulated Bits ({demod.get('bits_per_symbol', 0)} b/sym) | `{latencies.get('demodulation', 0)} ms` |")
        md.append(f"| **8** | **Signal Classification** | {sig_class.get('classified_standard', 'Standard Telemetry')} [{sig_class.get('threat_level', 'MONITORED')}] | `{latencies.get('signal_classification', 0)} ms` |\n")

        md.append("## 2. RF Parameters & Spectral Characteristics")
        md.append("| Parameter | Estimated Value | Unit |")
        md.append("| :--- | :--- | :--- |")
        md.append(f"| **Carrier Center Frequency (fc)** | {freq.get('carrier_freq_hz', 0)/1e6:.4f} | MHz |")
        md.append(f"| **Occupied Bandwidth (99% OBW)** | {freq.get('occupied_bandwidth_99_hz', freq.get('occupied_bw_99_hz', 0))/1e3:.2f} | kHz |")
        md.append(f"| **Symbol Rate (Rs)** | {freq.get('symbol_rate_baud', freq.get('baud_rate_hz', 0))/1e3:.2f} | kBaud |")
        md.append(f"| **Residual Carrier Offset (CFO)** | {freq.get('cfo_residual_hz', freq.get('cfo_estimated_hz', 0)):.2f} | Hz |")
        md.append(f"| **Samples Per Symbol (SPS)** | {freq.get('samples_per_symbol', 0):.2f} | samples/sym |")
        md.append(f"| **Consensus SNR** | {noise.get('snr_db', noise.get('snr_m2m4_db', 0)):.2f} | dB |")
        md.append(f"| **Noise Margin** | {noise.get('noise_margin_db', 0):.2f} | dB |\n")

        md.append("## 3. Modulation Recognition & Cumulants")
        cum = amc.get("cumulants", {})
        md.append(f"- **C20 (2nd Order):** `{cum.get('C20', 0):.4f}`")
        md.append(f"- **C40 (4th Order):** `{cum.get('C40', 0):.4f}`")
        md.append(f"- **C42 (4th Order Magnitude):** `{cum.get('C42', 0):.4f}`\n")
        
        probs = amc.get("probabilities", {})
        if probs:
            md.append("| Candidate Modulation | Soft Probability Score |")
            md.append("| :--- | :--- |")
            for mod_name, score in sorted(probs.items(), key=lambda x: x[1], reverse=True):
                bar = "█" * int(score * 20)
                md.append(f"| **{mod_name}** | `{score*100:.1f}%` {bar} |")
            md.append("")

        md.append("## 4. Recovered Intelligence Stream & Decoded Telemetry")
        md.append(f"- **Classified Standard:** `{sig_class.get('classified_standard', 'Standard Telemetry')}`")
        md.append(f"- **Emitter Category:** `{sig_class.get('emitter_category', 'General Transceiver')}`")
        md.append(f"- **Sync Marker:** `{framing.get('sync_marker', 'N/A')}` (Lock: `{framing.get('lock_status', 'LOCKED')}`)")
        md.append(f"- **CRC Verification:** `{crc.get('status', 'VALID')} [{crc.get('computed_crc', 'N/A')}]`")
        md.append(f"- **FEC Decoder:** `{fec.get('scheme', 'N/A')}` (Corrected Errors: `{fec.get('corrected_errors', 0)}`)")
        md.append(f"- **Shannon Entropy:** `{telem.get('shannon_entropy', 0):.2f} bits/byte`\n")
        md.append(f"```text\n{telem.get('decoded_text', 'NO_PAYLOAD_LOCKED')}\n```\n")
        md.append("### Raw Hex Byte Dump")
        md.append(f"```hex\n{telem.get('hex_dump', 'N/A')}\n```\n")

        md.append("---\n*AetherSpectra Autonomous Signal Intelligence Platform*")
        return "\n".join(md)

    @staticmethod
    def generate_csv(data: Dict[str, Any]) -> str:
        """Generates structured CSV report string."""
        output = io.StringIO()
        writer = csv.writer(output)

        writer.writerow(["AetherSpectra Signal Intelligence & Telemetry Analytics Dossier"])
        writer.writerow([])

        # Section 1: Metadata
        writer.writerow(["Section", "Parameter", "Value", "Unit / Notes"])
        writer.writerow(["Metadata", "Filename", data.get("filename", "N/A"), ""])
        writer.writerow(["Metadata", "File ID", data.get("file_id", "N/A"), ""])
        writer.writerow(["Metadata", "Total Latency", data.get("total_latency_ms", 0), "ms"])
        writer.writerow(["Metadata", "Sample Count", data.get("sample_count", 0), "samples"])

        # Section 2: 8-Stage Latencies
        latencies = data.get("latencies", {})
        for stage, lat in latencies.items():
            writer.writerow(["Latency Trace", stage.replace("_", " ").title(), lat, "ms"])

        # Section 3: Physical Parameters
        freq = data.get("frequency_estimation", {}) or data.get("parameters", {})
        noise = data.get("noise_estimation", {}) or data.get("preprocessing_stats", {})
        writer.writerow(["Parameters", "Carrier Frequency", freq.get("carrier_freq_hz", 0), "Hz"])
        writer.writerow(["Parameters", "Occupied 99% BW", freq.get("occupied_bandwidth_99_hz", freq.get("occupied_bw_99_hz", 0)), "Hz"])
        writer.writerow(["Parameters", "Symbol Rate", freq.get("symbol_rate_baud", freq.get("baud_rate_hz", 0)), "symbols/sec"])
        writer.writerow(["Parameters", "Residual CFO", freq.get("cfo_residual_hz", freq.get("cfo_estimated_hz", 0)), "Hz"])
        writer.writerow(["Parameters", "SNR Consensus", noise.get("snr_db", noise.get("snr_m2m4_db", 0)), "dB"])

        # Section 4: AMC
        amc = data.get("amc", {})
        writer.writerow(["AMC Classifier", "Identified Modulation", amc.get("identified_modulation", "N/A"), ""])
        writer.writerow(["AMC Classifier", "Confidence", amc.get("confidence", 0), "score [0-1]"])

        # Section 5: Classification & Telemetry
        sig_class = data.get("signal_classification", {})
        telem = data.get("telemetry", {})
        crc = data.get("crc", {})
        writer.writerow(["Classification", "Standard", sig_class.get("classified_standard", "N/A"), ""])
        writer.writerow(["Classification", "Emitter Category", sig_class.get("emitter_category", "N/A"), ""])
        writer.writerow(["Classification", "CRC Check", crc.get("status", "N/A"), ""])
        writer.writerow(["Classification", "Decoded Telemetry", telem.get("decoded_text", "N/A"), "ASCII"])

        return output.getvalue()

    @staticmethod
    def generate_docx(data: Dict[str, Any]) -> bytes:
        """Generates formatted Microsoft Word document (.docx)."""
        doc = Document()

        style = doc.styles['Normal']
        font = style.font
        font.name = 'Calibri'
        font.size = Pt(11)
        font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

        # Title
        title_p = doc.add_paragraph()
        title_run = title_p.add_run("AETHER SPECTRA DEFENSE & RF SYSTEMS")
        title_run.font.size = Pt(16)
        title_run.font.bold = True
        title_run.font.color.rgb = RGBColor(0x02, 0x84, 0xC7)
        title_p.paragraph_format.space_after = Pt(2)

        subtitle_p = doc.add_paragraph()
        subtitle_run = subtitle_p.add_run("Autonomous Signal Intelligence & Telemetry Analytics Dossier")
        subtitle_run.font.size = Pt(13)
        subtitle_run.font.bold = True
        subtitle_run.font.color.rgb = RGBColor(0x37, 0x41, 0x51)
        subtitle_p.paragraph_format.space_after = Pt(12)

        # Metadata Table
        filename = data.get("filename", "Unknown")
        file_id = data.get("file_id", "N/A")
        total_lat = data.get("total_latency_ms", 0)
        
        meta_table = doc.add_table(rows=3, cols=2)
        meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        meta_table.rows[0].cells[0].text = "Target File Name"
        meta_table.rows[0].cells[1].text = str(filename)
        meta_table.rows[1].cells[0].text = "Signal Identifier"
        meta_table.rows[1].cells[1].text = str(file_id)
        meta_table.rows[2].cells[0].text = "End-to-End Processing Latency"
        meta_table.rows[2].cells[1].text = f"{total_lat} ms"
        for r in meta_table.rows:
            r.cells[0].paragraphs[0].runs[0].font.bold = True
        doc.add_paragraph().paragraph_format.space_after = Pt(8)

        # 8-Stage DSP Table
        h1 = doc.add_heading("1. 8-Stage DSP Pipeline Execution Summary", level=2)
        h1.paragraph_format.space_before = Pt(10)
        
        latencies = data.get("latencies", {})
        p_table = doc.add_table(rows=9, cols=3)
        p_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        headers = ["Stage #", "DSP Stage Name", "Latency (ms)"]
        for i, h in enumerate(headers):
            p_table.rows[0].cells[i].text = h
            p_table.rows[0].cells[i].paragraphs[0].runs[0].font.bold = True

        stages_list = [
            ("1", "Filtering (Bandpass/DC/IQ Balance)", f"{latencies.get('filtering', 0)} ms"),
            ("2", "Fast Fourier Transform (FFT & PSD)", f"{latencies.get('fft', 0)} ms"),
            ("3", "Frequency Estimation (fc, BW, CFO)", f"{latencies.get('frequency_estimation', 0)} ms"),
            ("4", "Synchronization (Carrier & Timing STR)", f"{latencies.get('synchronization', 0)} ms"),
            ("5", "Modulation Detection (AMC & Cumulants)", f"{latencies.get('modulation_detection', 0)} ms"),
            ("6", "Noise Estimation (M2M4, N0, Margin)", f"{latencies.get('noise_estimation', 0)} ms"),
            ("7", "Demodulation (Symbols → Bits LLR)", f"{latencies.get('demodulation', 0)} ms"),
            ("8", "Signal Classification & Intelligence", f"{latencies.get('signal_classification', 0)} ms")
        ]
        for row_idx, (num, name, lat) in enumerate(stages_list, start=1):
            p_table.rows[row_idx].cells[0].text = num
            p_table.rows[row_idx].cells[1].text = name
            p_table.rows[row_idx].cells[2].text = lat
        doc.add_paragraph().paragraph_format.space_after = Pt(8)

        # Telemetry Payload
        sig_class = data.get("signal_classification", {})
        telem = data.get("telemetry", {})
        h2 = doc.add_heading("2. Decoded Telemetry Payload & Intelligence", level=2)
        h2.paragraph_format.space_before = Pt(10)

        p_info = doc.add_paragraph()
        p_info.add_run("Classified Standard: ").font.bold = True
        p_info.add_run(f"{sig_class.get('classified_standard', 'Standard Telemetry')}\n")
        p_info.add_run("Emitter Category: ").font.bold = True
        p_info.add_run(f"{sig_class.get('emitter_category', 'General Transceiver')}\n")

        tel_p = doc.add_paragraph()
        tel_p.add_run("Decoded ASCII Stream:\n").font.bold = True
        tel_box = doc.add_paragraph()
        tel_box_run = tel_box.add_run(telem.get('decoded_text', 'NO_TELEMETRY_LOCKED'))
        tel_box_run.font.name = 'Courier New'
        tel_box_run.font.size = Pt(10)
        tel_box_run.font.bold = True
        tel_box_run.font.color.rgb = RGBColor(0x03, 0x69, 0xA1)

        doc_io = io.BytesIO()
        doc.save(doc_io)
        return doc_io.getvalue()

    @staticmethod
    def generate_pdf(data: Dict[str, Any]) -> bytes:
        """Generates high-quality PDF report using ReportLab."""
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            rightMargin=36,
            leftMargin=36,
            topMargin=36,
            bottomMargin=36
        )

        styles = getSampleStyleSheet()
        
        title_style = ParagraphStyle(
            'DocTitle',
            parent=styles['Heading1'],
            fontName='Helvetica-Bold',
            fontSize=16,
            leading=20,
            textColor=colors.HexColor('#0284C7')
        )
        subtitle_style = ParagraphStyle(
            'DocSubTitle',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=11,
            leading=14,
            textColor=colors.HexColor('#475569')
        )
        section_heading = ParagraphStyle(
            'SectionHeading',
            parent=styles['Heading2'],
            fontName='Helvetica-Bold',
            fontSize=12,
            leading=16,
            textColor=colors.HexColor('#0F172A'),
            spaceBefore=10,
            spaceAfter=6
        )
        normal_style = ParagraphStyle(
            'Body',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=9,
            leading=12,
            textColor=colors.HexColor('#1E293B')
        )
        mono_style = ParagraphStyle(
            'Mono',
            parent=styles['Normal'],
            fontName='Courier-Bold',
            fontSize=9,
            leading=12,
            textColor=colors.HexColor('#0369A1')
        )
        mono_small = ParagraphStyle(
            'MonoSmall',
            parent=styles['Normal'],
            fontName='Courier',
            fontSize=8,
            leading=10,
            textColor=colors.HexColor('#334155')
        )

        story = []

        # Header Title
        story.append(Paragraph("AETHER SPECTRA DEFENSE & RF SYSTEMS", title_style))
        story.append(Paragraph("Autonomous Signal Intelligence & Telemetry Analytics Dossier", subtitle_style))
        story.append(Spacer(1, 8))
        story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#0284C7'), spaceBefore=2, spaceAfter=8))

        # Metadata Table
        filename = data.get("filename", "Unknown")
        file_id = data.get("file_id", "N/A")
        total_lat = data.get("total_latency_ms", 0)
        amc = data.get("amc", {})
        noise = data.get("noise_estimation", {}) or data.get("preprocessing_stats", {})
        sync = data.get("synchronization", {})
        sig_class = data.get("signal_classification", {})

        meta_data = [
            [Paragraph("<b>Target File:</b>", normal_style), Paragraph(str(filename), normal_style),
             Paragraph("<b>Processing Latency:</b>", normal_style), Paragraph(f"{total_lat} ms", normal_style)],
            [Paragraph("<b>File ID:</b>", normal_style), Paragraph(str(file_id)[:24] + "...", mono_style),
             Paragraph("<b>Carrier Lock:</b>", normal_style), Paragraph(f"{sync.get('carrier_lock_status', 'N/A')}", normal_style)],
            [Paragraph("<b>Identified Modulation:</b>", normal_style), Paragraph(f"<b>{amc.get('identified_modulation', 'N/A')}</b> ({round(amc.get('confidence', 0)*100, 1)}%)", normal_style),
             Paragraph("<b>SNR:</b>", normal_style), Paragraph(f"{noise.get('snr_db', noise.get('snr_m2m4_db', 0)):.1f} dB", normal_style)]
        ]
        meta_table = Table(meta_data, colWidths=[1.3*inch, 2.3*inch, 1.4*inch, 2.0*inch])
        meta_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F8FAFC')),
            ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#E2E8F0')),
            ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        story.append(meta_table)
        story.append(Spacer(1, 10))

        # Section 1: 8-Stage DSP Execution Trace
        story.append(Paragraph("1. 8-Stage DSP Pipeline Execution Trace", section_heading))
        latencies = data.get("latencies", {})
        stage_names = [
            ("1. Filtering (Bandpass/DC/IQ Balance)", latencies.get("filtering", 0)),
            ("2. Fast Fourier Transform (FFT & PSD)", latencies.get("fft", 0)),
            ("3. Frequency Estimation (fc, BW, CFO)", latencies.get("frequency_estimation", 0)),
            ("4. Synchronization (Carrier & Timing STR)", latencies.get("synchronization", 0)),
            ("5. Modulation Detection (AMC & Cumulants)", latencies.get("modulation_detection", 0)),
            ("6. Noise Estimation (M2M4, N0, Margin)", latencies.get("noise_estimation", 0)),
            ("7. Demodulation (Symbols → Bits LLR)", latencies.get("demodulation", 0)),
            ("8. Signal Classification & Intelligence", latencies.get("signal_classification", 0))
        ]
        lat_rows = [[Paragraph("<b>DSP Pipeline Stage</b>", normal_style), Paragraph("<b>Execution Latency</b>", normal_style)]]
        for name, l in stage_names:
            lat_rows.append([Paragraph(name, normal_style), Paragraph(f"{l} ms", normal_style)])
        lat_rows.append([Paragraph("<b>Total End-to-End Latency</b>", normal_style), Paragraph(f"<b>{total_lat} ms</b>", normal_style)])

        lat_table = Table(lat_rows, colWidths=[4.8*inch, 2.2*inch])
        lat_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#E2E8F0')),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
            ('TOPPADDING', (0, 0), (-1, -1), 2.5),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 2.5),
        ]))
        story.append(lat_table)
        story.append(Spacer(1, 10))

        # Section 2: Intelligence & Decoded Telemetry
        story.append(Paragraph("2. Decoded Intelligence & Telemetry Payload", section_heading))
        telem = data.get("telemetry", {})
        telemetry_text = telem.get('decoded_text', 'NO_TELEMETRY_LOCKED')
        framing = data.get("framing", {})
        crc = data.get("crc", {})
        fec = data.get("fec", {})
        
        telemetry_data = [
            [Paragraph(f"<b>Decoded ASCII Stream:</b><br/>{telemetry_text}", mono_style)],
            [Paragraph(f"<b>Standard:</b> {sig_class.get('classified_standard', 'N/A')} | <b>Sync Lock:</b> {framing.get('lock_status', 'LOCKED')} | <b>CRC:</b> {crc.get('status', 'VALID')}", normal_style)],
            [Paragraph(f"<b>Raw Hex Stream:</b><br/>{telem.get('hex_dump', 'N/A')}", mono_small)]
        ]
        telemetry_table = Table(telemetry_data, colWidths=[7.0*inch])
        telemetry_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F1F5F9')),
            ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#94A3B8')),
            ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
            ('TOPPADDING', (0, 0), (-1, -1), 5),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ]))
        story.append(telemetry_table)

        doc.build(story)
        return buffer.getvalue()
